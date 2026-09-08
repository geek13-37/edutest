from __future__ import annotations

import uuid
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    Assignment,
    Attempt,
    AttemptStatus,
    Class,
    ClassMember,
    Question,
    School,
    Test,
    User,
    UserRole,
)
from app.services import class_service, grading
from app.services.results_service import assignment_results

STATUS_RU = {
    "submitted": "Сдан",
    "expired": "Время вышло",
    "in_progress": "В процессе",
    None: "Не приступал(а)",
}
FINISHED = (AttemptStatus.submitted, AttemptStatus.expired)


def _fill_sheet(ws: Worksheet, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width


def _fmt(value) -> str:
    return value.strftime("%d.%m.%Y %H:%M") if value is not None else ""


def export_assignment_results_xlsx(db: Session, teacher: User, assignment_id: uuid.UUID) -> BytesIO:
    data = assignment_results(db, teacher, assignment_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"

    rows = [
        [
            idx,
            r["student_name"],
            r["student_login"],
            r["attempts_used"],
            r["best_score"] if r["best_score"] is not None else "",
            r["max_score"] if r["max_score"] is not None else "",
            f'{r["best_percent"]}%' if r["best_percent"] is not None else "",
            r["grade"] or "",
            STATUS_RU.get(r["status"].value if r["status"] else None, ""),
        ]
        for idx, r in enumerate(data["rows"], start=1)
    ]
    _fill_sheet(
        ws,
        ["№", "ФИО", "Логин", "Попыток", "Баллы", "Макс. балл", "Процент", "Оценка", "Статус"],
        rows,
        [5, 28, 28, 9, 8, 11, 10, 8, 14],
    )

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def export_school_xlsx(db: Session, school_id: uuid.UUID) -> BytesIO:
    """Полная выгрузка школы: одна книга, отдельный лист на каждую сущность."""
    school = db.get(School, school_id)
    if school is None:
        raise ValueError("Школа не найдена")

    teachers = list(
        db.scalars(
            select(User)
            .where(User.school_id == school_id, User.role == UserRole.teacher)
            .order_by(User.full_name)
        )
    )
    students = list(
        db.scalars(
            select(User)
            .where(User.school_id == school_id, User.role == UserRole.student)
            .order_by(User.full_name)
        )
    )
    classes = list(
        db.scalars(
            select(Class).where(Class.school_id == school_id).order_by(Class.created_at)
        )
    )
    creators = {u.id: u.full_name for u in db.scalars(select(User).where(User.role == UserRole.teacher))}
    members: dict[uuid.UUID, list[str]] = {}
    for cm, klass, student in db.execute(
        select(ClassMember, Class, User)
        .join(Class, Class.id == ClassMember.class_id)
        .join(User, User.id == ClassMember.student_id)
        .where(Class.school_id == school_id)
    ):
        members.setdefault(student.id, []).append(class_service.display_name(klass))

    tests = list(
        db.scalars(
            select(Test)
            .join(User, User.id == Test.owner_id)
            .where(User.school_id == school_id)
            .order_by(Test.created_at)
        )
    )
    test_owner = {t.id: creators.get(t.owner_id, "") for t in tests}

    wb = Workbook()

    ws = wb.active
    ws.title = "Школа"
    _fill_sheet(
        ws,
        ["Параметр", "Значение"],
        [
            ["Название", school.name],
            ["Город", school.city],
            ["Регион", school.region or ""],
            ["Код регистрации", school.signup_code],
            ["Префикс логинов", school.login_prefix],
            ["Создана", _fmt(school.created_at)],
            ["Статус", "В архиве" if school.archived_at else "Активна"],
            ["Учителей", len(teachers)],
            ["Классов", len(classes)],
            ["Учеников", len(students)],
        ],
        [22, 46],
    )

    _fill_sheet(
        wb.create_sheet("Учителя"),
        ["ФИО", "Email", "Статус", "Создан"],
        [
            [t.full_name, t.email or "", "Активен" if t.is_active else "Отключен", _fmt(t.created_at)]
            for t in teachers
        ],
        [30, 30, 12, 18],
    )

    _fill_sheet(
        wb.create_sheet("Классы"),
        ["Класс", "Создал", "Учеников", "В архиве", "Создан"],
        [
            [
                class_service.display_name(c),
                creators.get(c.teacher_id, ""),
                db.scalar(
                    select(func.count(ClassMember.id)).where(ClassMember.class_id == c.id)
                ) or 0,
                "да" if c.archived else "нет",
                _fmt(c.created_at),
            ]
            for c in classes
        ],
        [18, 28, 10, 10, 18],
    )

    _fill_sheet(
        wb.create_sheet("Ученики"),
        ["ФИО", "Логин", "Классы", "Создан"],
        [
            [s.full_name, s.username or "", ", ".join(sorted(members.get(s.id, []))), _fmt(s.created_at)]
            for s in students
        ],
        [30, 14, 24, 18],
    )

    _fill_sheet(
        wb.create_sheet("Тесты"),
        ["Название", "Автор", "Статус", "Вопросов", "Создан", "Обновлен"],
        [
            [
                t.title,
                test_owner.get(t.id, ""),
                "Опубликован" if t.status.value == "published" else "Черновик",
                db.scalar(select(func.count(Question.id)).where(Question.test_id == t.id)) or 0,
                _fmt(t.created_at),
                _fmt(t.updated_at),
            ]
            for t in tests
        ],
        [34, 26, 14, 10, 18, 18],
    )

    # ── назначения и результаты ──────────────────────────
    assignments = list(
        db.scalars(
            select(Assignment)
            .join(Class, Class.id == Assignment.class_id)
            .where(Class.school_id == school_id)
            .order_by(Assignment.created_at)
        )
    )
    assign_rows = []
    result_rows = []
    for a in assignments:
        test = db.get(Test, a.test_id)
        klass = db.get(Class, a.class_id)
        class_name = class_service.display_name(klass) if klass else ""
        test_title = test.title if test else ""

        roster = list(
            db.scalars(
                select(User)
                .join(ClassMember, ClassMember.student_id == User.id)
                .where(ClassMember.class_id == a.class_id)
                .order_by(User.full_name)
            )
        )
        attempts = list(
            db.scalars(select(Attempt).where(Attempt.assignment_id == a.id))
        )
        by_student: dict[uuid.UUID, list[Attempt]] = {}
        for at in attempts:
            by_student.setdefault(at.student_id, []).append(at)

        percents: list[float] = []
        submitted = 0
        for s in roster:
            finished = [x for x in by_student.get(s.id, []) if x.status in FINISHED]
            best = max(finished, key=lambda x: x.percent, default=None)
            if best is not None:
                submitted += 1
                percents.append(best.percent)
            result_rows.append(
                [
                    test_title,
                    class_name,
                    s.full_name,
                    s.username or "",
                    len(finished),
                    f"{best.percent}%" if best else "",
                    grading.grade_letter(best.percent, test.grade_thresholds) if best and test else "",
                    STATUS_RU.get(best.status.value if best else None, "Не приступал(а)"),
                ]
            )
        avg = round(sum(percents) / len(percents), 1) if percents else ""
        assign_rows.append(
            [
                test_title,
                class_name,
                _fmt(a.opens_at),
                _fmt(a.closes_at),
                a.max_attempts if a.max_attempts is not None else "без лимита",
                f"{submitted} / {len(roster)}",
                f"{avg}%" if avg != "" else "",
            ]
        )

    _fill_sheet(
        wb.create_sheet("Назначения"),
        ["Тест", "Класс", "Открыт", "Закрыт", "Попыток", "Сдали", "Средний %"],
        assign_rows,
        [30, 14, 18, 18, 12, 12, 12],
    )
    _fill_sheet(
        wb.create_sheet("Результаты"),
        ["Тест", "Класс", "ФИО", "Логин", "Попыток", "Процент", "Оценка", "Статус"],
        result_rows,
        [28, 12, 28, 14, 10, 10, 8, 16],
    )

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
